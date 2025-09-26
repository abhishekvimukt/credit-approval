import openpyxl
from celery import shared_task
from .models import Customer, Loan
from django.db import transaction
from datetime import datetime

@shared_task
def ingest_customer_data():
    try:
        # Assuming customer_data.xlsx is in the project root
        workbook = openpyxl.load_workbook('customer_data.xlsx')
        sheet = workbook.active

        with transaction.atomic():
            for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True)):
                # Skip empty rows
                if not any(row):
                    continue
                
                # Extract data
                customer_id = row[0]
                first_name = row[1]
                last_name = row[2]
                age = row[3]
                monthly_salary = row[4]
                phone_number = row[5]

                # Calculate approved_limit: 36 * monthly_salary, rounded to nearest lakh
                calculated_limit = 36 * monthly_salary
                approved_limit = (round(calculated_limit / 100000) * 100000)

                Customer.objects.update_or_create(
                    customer_id=customer_id,
                    defaults={
                        'first_name': first_name,
                        'last_name': last_name,
                        'age': age,
                        'monthly_salary': monthly_salary,
                        'approved_limit': approved_limit,
                        'phone_number': phone_number
                    }
                )
        return "Customer data ingestion completed successfully."
    except FileNotFoundError:
        return "Error: customer_data.xlsx not found."
    except Exception as e:
        return f"Error ingesting customer data: {e}"


@shared_task
def ingest_loan_data():
    try:
        # Assuming loan_data.xlsx is in the project root
        workbook = openpyxl.load_workbook('loan_data.xlsx')
        sheet = workbook.active

        with transaction.atomic():
            for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True)):
                # Skip empty rows
                if not any(row):
                    continue

                # Extract data
                customer_id = row[0]
                loan_id = row[1]
                loan_amount = row[2]
                tenure = row[3]
                interest_rate = row[4]
                monthly_installment = row[5]
                date_approved_raw = row[6]
                end_date_raw = row[7]

                date_approved = None
                if isinstance(date_approved_raw, datetime):
                    date_approved = date_approved_raw.date()
                elif isinstance(date_approved_raw, (int, float)):
                    # Convert Excel serial date to datetime.date
                    # Excel epoch is 1899-12-30. Unix epoch is 1970-01-01.
                    # 25569 is the difference in days between these two epochs.
                    # 86400 is seconds in a day.
                    try:
                        date_approved = datetime.fromtimestamp((date_approved_raw - 25569) * 86400).date()
                    except (ValueError, OSError): # Handle potential out-of-range timestamps
                        date_approved = None
                else:
                    # Try to convert to string and then parse
                    date_str = str(date_approved_raw)
                    for fmt in ('%Y-%m-%d', '%m/%d/%Y', '%d-%m-%Y', '%d/%m/%Y'): # Add more common formats if needed
                        try:
                            date_approved = datetime.strptime(date_str, fmt).date()
                            break
                        except ValueError:
                            continue
                    if date_approved is None: # If loop finishes without successful parse
                        print(f"Warning: Could not parse date_approved '{date_approved_raw}'")

                end_date = None
                if isinstance(end_date_raw, datetime):
                    end_date = end_date_raw.date()
                elif isinstance(end_date_raw, (int, float)):
                    try:
                        end_date = datetime.fromtimestamp((end_date_raw - 25569) * 86400).date()
                    except (ValueError, OSError):
                        end_date = None
                else:
                    # Try to convert to string and then parse
                    date_str = str(end_date_raw)
                    for fmt in ('%Y-%m-%d', '%m/%d/%Y', '%d-%m-%Y', '%d/%m/%Y'):
                        try:
                            end_date = datetime.strptime(date_str, fmt).date()
                            break
                        except ValueError:
                            continue
                    if end_date is None:
                        print(f"Warning: Could not parse end_date '{end_date_raw}'")

                try:
                    customer = Customer.objects.get(customer_id=customer_id)
                except Customer.DoesNotExist:
                    print(f"Skipping loan {loan_id}: Customer {customer_id} not found.")
                    continue

                Loan.objects.update_or_create(
                    loan_id=loan_id,
                    defaults={
                        'customer': customer,
                        'loan_amount': loan_amount,
                        'tenure': tenure,
                        'interest_rate': interest_rate,
                        'monthly_installment': monthly_installment,
                        'date_approved': date_approved,
                        'end_date': end_date,
                        'emis_paid_on_time': 0, # Assuming initial ingestion, EMIs paid will be 0
                        'num_of_repayments': 0, # Assuming initial ingestion, repayments will be 0
                        'loan_status': 'Approved' # Assuming historical loans are approved
                    }
                )
        return "Loan data ingestion completed successfully."
    except FileNotFoundError:
        return "Error: loan_data.xlsx not found."
    except Exception as e:
        return f"Error ingesting loan data: {e}"
